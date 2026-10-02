from __future__ import annotations

import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HEADERS = [
    "Name",
    "Offer",
    "Description",
    "Category",
    "Subcategory",
    "Location",
    "Tags",
    "GitHub pack",
    "Underrated",
    "Verification",
    "Original link",
    "Claim link",
    "Logo",
    "Synced at",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Arial", size=10)
LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")
WRAP = Alignment(vertical="top", wrap_text=True)


def _tags(raw: str | None) -> str:
    try:
        tags = json.loads(raw or "[]")
    except json.JSONDecodeError:
        return ""
    return ", ".join(tags)


def build_workbook(rows) -> BytesIO:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Offers"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        original = row["canonical_url"]
        claim = row["claim_url"]
        sheet.append([
            row["name"],
            row["offer"],
            row["description"] or "",
            row["category_main"] or "",
            row["category_sub"] or "",
            row["location"] or "",
            _tags(row["tags_json"]),
            "Yes" if row["github_offer"] else "No",
            "Yes" if row["is_underrated"] else "No",
            row["verification_status"] or "",
            original or "",
            claim or "",
            row["logo"] or "",
            row["synced_at"] or "",
        ])
        excel_row = sheet.max_row
        for col in (11, 12):
            cell = sheet.cell(excel_row, col)
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = LINK_FONT
        for col in range(1, 15):
            if col not in (11, 12):
                sheet.cell(excel_row, col).font = BODY_FONT
            sheet.cell(excel_row, col).alignment = WRAP

    widths = [28, 42, 46, 24, 22, 18, 22, 14, 14, 16, 46, 46, 36, 24]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.auto_filter.ref = f"A1:N{sheet.max_row}"
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 22
    if sheet.max_row > 1:
        table = Table(displayName="Offers", ref=f"A1:N{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)

    summary = wb.create_sheet("Summary", 0)
    summary["A1"] = "StudentOffers.co directory export"
    summary["A1"].font = Font(name="Arial", bold=True, size=16, color="1F4E79")
    summary["A2"] = "Original link is the StudentOffers page. Claim link is the brand page. Re-check both before using a deal."
    summary["A2"].font = Font(name="Arial", size=10, italic=True)
    summary["A4"] = "Offers"
    summary["B4"] = f"=COUNTA(Offers!A2:A{sheet.max_row})"
    summary["A5"] = "GitHub pack"
    summary["B5"] = f'=COUNTIF(Offers!H2:H{sheet.max_row},"Yes")'
    summary["A6"] = "Source"
    summary["B6"] = "https://www.studentoffers.co/api/offers"
    summary["B6"].hyperlink = "https://www.studentoffers.co/api/offers"
    summary["B6"].font = LINK_FONT
    for cell in ("A4", "A5", "A6"):
        summary[cell].font = Font(name="Arial", bold=True, size=11)
    summary["B4"].font = Font(name="Arial", size=11)
    summary["B5"].font = Font(name="Arial", size=11)
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 78

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
